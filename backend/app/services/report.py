"""
Monta o relatório em Excel (.xlsx) a partir de transações já filtradas.

Abas:
- Resumo: filtros, números do período (cada um com "o que significa"),
  destaques, os maiores gastos e um guia de leitura
- Por mês: entradas, saídas, saldo, saldo acumulado e economia de cada mês,
  com gráfico
- Por categoria: saídas e entradas por categoria, com % e média, e gráfico
- Transações: a lista completa (filtro do Excel e cabeçalho fixo)

Nos resumos, "Saídas" aparece como valor positivo (quanto saiu); o sinal só
fica na lista de transações, igual ao extrato do banco.

Transações de categorias marcadas como "ignorar nos gráficos" (ex:
pagamento de fatura) aparecem na aba Transações, mas ficam fora de todos
os totais, igual ao dashboard.
"""
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.transaction import Transaction

# ---------- Visual (cores da Vexira) ----------

FONT_NAME = "Arial"
NAVY = "0B1F33"
CYAN = "0891B2"  # ciano mais escuro que o do app: legível em fundo branco
GREEN = "1A7F55"
RED = "C0392B"
MUTED = "5B6B7B"
ZEBRA = "F3F7FA"
LINE = "D9E2EC"

MONEY_FORMAT = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
# Na lista de transações: entrada em verde com "+", saída em vermelho com "-"
SIGNED_MONEY_FORMAT = '[Color10]+"R$" #,##0.00;[Red]-"R$" #,##0.00;"R$" 0.00'
PERCENT_FORMAT = "0.0%"
DATE_FORMAT = "DD/MM/YYYY"
COUNT_FORMAT = "0"

UNCATEGORIZED = "Sem categoria"
MONTHS = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

# Tabelas começam na linha 4: 1 = título, 2 = explicação, 3 = respiro
TABLE_START = 4


def _font(**kwargs) -> Font:
    kwargs.setdefault("size", 10)
    return Font(name=FONT_NAME, **kwargs)


HEADER_FONT = _font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
BOTTOM_LINE = Border(bottom=Side(style="thin", color=LINE))
TOTAL_LINE = Border(top=Side(style="medium", color=NAVY))
WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass
class ReportFilters:
    start_date: Optional[date] = None
    end_date: Optional[date] = None
    statement_name: Optional[str] = None


def counts_in_reports(transaction: Transaction) -> bool:
    return not (transaction.category and transaction.category.ignore_in_reports)


def _category_name(transaction: Transaction) -> str:
    return transaction.category.name if transaction.category else UNCATEGORIZED


def month_label(year: int, month: int) -> str:
    """(2025, 3) -> "mar/2025" """
    return f"{MONTHS[month - 1]}/{year}"


# ---------- Números do período ----------


@dataclass
class Totals:
    """Tudo o que as abas mostram, calculado uma vez só."""

    counted: list[Transaction]
    total_count: int
    income: Decimal = Decimal(0)
    expenses: Decimal = Decimal(0)  # positivo: quanto saiu
    # (ano, mês) -> [entradas, saídas, quantidade]
    months: dict = field(default_factory=dict)
    # categoria -> [valor, quantidade]
    spending: dict = field(default_factory=dict)
    earning: dict = field(default_factory=dict)

    @property
    def balance(self) -> Decimal:
        return self.income - self.expenses


def _totals(transactions: list[Transaction]) -> Totals:
    counted = [t for t in transactions if counts_in_reports(t)]
    totals = Totals(counted=counted, total_count=len(transactions))
    months: dict = defaultdict(lambda: [Decimal(0), Decimal(0), 0])
    spending: dict = defaultdict(lambda: [Decimal(0), 0])
    earning: dict = defaultdict(lambda: [Decimal(0), 0])

    for t in counted:
        month = months[(t.date.year, t.date.month)]
        month[2] += 1
        if t.amount > 0:
            totals.income += t.amount
            month[0] += t.amount
            earning[_category_name(t)][0] += t.amount
            earning[_category_name(t)][1] += 1
        elif t.amount < 0:
            totals.expenses += -t.amount
            month[1] += -t.amount
            spending[_category_name(t)][0] += -t.amount
            spending[_category_name(t)][1] += 1

    totals.months = dict(sorted(months.items()))
    # Maior valor primeiro; empate em ordem alfabética
    totals.spending = dict(sorted(spending.items(), key=lambda i: (-i[1][0], i[0])))
    totals.earning = dict(sorted(earning.items(), key=lambda i: (-i[1][0], i[0])))
    return totals


def build_report(
    transactions: list[Transaction], filters: ReportFilters, generated_at: datetime
) -> bytes:
    transactions = sorted(transactions, key=lambda t: (t.date, t.id))
    totals = _totals(transactions)
    ignored = [t for t in transactions if not counts_in_reports(t)]

    workbook = Workbook()
    _summary_sheet(workbook.active, totals, ignored, filters, generated_at)
    _monthly_sheet(workbook.create_sheet("Por mês"), totals)
    _category_sheet(workbook.create_sheet("Por categoria"), totals)
    _transactions_sheet(workbook.create_sheet("Transações"), transactions)

    for sheet in workbook.worksheets:
        _finish_sheet(sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------- Aba Resumo ----------


def _summary_sheet(sheet: Worksheet, totals: Totals, ignored, filters, generated_at):
    sheet.title = "Resumo"
    sheet.sheet_properties.tabColor = NAVY
    _set_widths(sheet, [34, 20, 70])
    _title(
        sheet,
        "Vexira — Relatório financeiro",
        "Controle financeiro inteligente. Os números abaixo resumem o período escolhido; "
        "as outras abas mostram os detalhes.",
        columns=3,
    )

    row = TABLE_START
    row = _section(sheet, row, "Filtros deste relatório")
    row = _label_rows(
        sheet,
        row,
        [
            ("Período", _period_label(filters), None, "Datas das transações incluídas"),
            ("Extrato", filters.statement_name or "Todos", None, "Arquivo importado considerado"),
            ("Gerado em", generated_at.strftime("%d/%m/%Y %H:%M"), None, "Data e hora desta exportação"),
            ("Transações", totals.total_count, COUNT_FORMAT, "Quantidade de lançamentos no período"),
        ],
    )
    if ignored:
        ignored_value = sum((abs(t.amount) for t in ignored), Decimal(0))
        row = _label_rows(
            sheet,
            row,
            [
                (
                    "Fora dos totais (categorias ignoradas)",
                    len(ignored),
                    COUNT_FORMAT,
                    f"{_money_text(ignored_value)} em categorias marcadas como \"ignorar nos gráficos\" "
                    "(ex: pagamento de fatura, investimentos). Estão na aba Transações, mas não "
                    "entram em nenhum total, para não contar o mesmo gasto duas vezes.",
                )
            ],
        )

    row = _section(sheet, row + 1, "Resumo do período")
    savings = totals.balance / totals.income if totals.income else None
    month_count = len(totals.months)
    row = _label_rows(
        sheet,
        row,
        [
            ("Entradas", totals.income, MONEY_FORMAT, "Todo o dinheiro que entrou: salário, Pix e transferências recebidas, estornos..."),
            ("Saídas", totals.expenses, MONEY_FORMAT, "Todo o dinheiro que saiu: compras, contas, Pix e transferências enviadas..."),
            ("Saldo", totals.balance, MONEY_FORMAT, "Entradas menos saídas. Positivo: sobrou dinheiro; negativo: gastou mais do que entrou."),
            (
                "Taxa de economia",
                float(savings) if savings is not None else "—",
                PERCENT_FORMAT,
                "Quanto das entradas sobrou (saldo ÷ entradas). Quanto maior, melhor.",
            ),
            (
                "Média de saídas por mês",
                totals.expenses / month_count if month_count else Decimal(0),
                MONEY_FORMAT,
                f"Saídas divididas por {_plural(month_count, 'mês', 'meses')} com movimentação no período.",
            ),
        ],
        colors={"Entradas": GREEN, "Saídas": RED, "Saldo": GREEN if totals.balance >= 0 else RED},
    )

    row = _section(sheet, row + 1, "Destaques")
    row = _label_rows(sheet, row, _highlights(totals))

    row = _section(sheet, row + 1, "Os 5 maiores gastos")
    biggest = sorted((t for t in totals.counted if t.amount < 0), key=lambda t: (t.amount, t.date))[:5]
    if biggest:
        row = _header_row(sheet, row, ["Descrição", "Valor", "Categoria e data"])
        for index, t in enumerate(biggest):
            _write_row(
                sheet,
                row,
                [t.description, -t.amount, f"{_category_name(t)} · {t.date.strftime('%d/%m/%Y')}"],
                formats={2: MONEY_FORMAT},
                zebra=index % 2 == 1,
            )
            sheet.cell(row=row, column=2).font = _font(color=RED)
            row += 1
    else:
        _note(sheet, row, "Nenhum gasto no período.", columns=3)
        row += 1

    row = _section(sheet, row + 1, "Como ler este relatório")
    guide = [
        "• Por mês: entradas, saídas e saldo de cada mês, o saldo acumulado e quanto sobrou (%). O gráfico compara entradas e saídas.",
        "• Por categoria: para onde foi o dinheiro (saídas) e de onde veio (entradas), com a fatia de cada categoria.",
        "• Transações: todos os lançamentos, com filtro no cabeçalho. Entradas em verde (+), saídas em vermelho (−).",
        "• \"Sem categoria\" reúne o que nenhuma regra reconheceu. Categorize na tela Transações da Vexira para relatórios mais precisos.",
    ]
    for line in guide:
        _note(sheet, row, line, columns=3)
        row += 1

    sheet.sheet_view.showGridLines = False


def _highlights(totals: Totals) -> list[tuple]:
    """Linhas de "Destaques": maior gasto, maior entrada, categoria e mês que mais pesaram."""
    expenses = [t for t in totals.counted if t.amount < 0]
    incomes = [t for t in totals.counted if t.amount > 0]
    rows = []

    if expenses:
        biggest = min(expenses, key=lambda t: (t.amount, t.date))
        rows.append(("Maior gasto", -biggest.amount, MONEY_FORMAT, f"{biggest.description} em {biggest.date.strftime('%d/%m/%Y')}"))
    else:
        rows.append(("Maior gasto", "—", None, "Nenhum gasto no período"))

    if incomes:
        biggest = max(incomes, key=lambda t: (t.amount, t.date))
        rows.append(("Maior entrada", biggest.amount, MONEY_FORMAT, f"{biggest.description} em {biggest.date.strftime('%d/%m/%Y')}"))
    else:
        rows.append(("Maior entrada", "—", None, "Nenhuma entrada no período"))

    # "Sem categoria" tem linha própria logo abaixo
    categorized = [(n, v) for n, v in totals.spending.items() if n != UNCATEGORIZED]
    if categorized:
        name, (value, count) = categorized[0]
        share = value / totals.expenses
        rows.append(
            (
                "Categoria que mais pesou",
                name,
                None,
                f"{_money_text(value)} em {_plural(count, 'transação', 'transações')}: "
                f"{_percent_text(share)} das saídas",
            )
        )
    else:
        rows.append(("Categoria que mais pesou", "—", None, "Nenhum gasto categorizado no período"))

    if totals.months and totals.expenses:
        (year, month), (_income, spent, _count) = max(totals.months.items(), key=lambda i: (i[1][1], i[0]))
        rows.append(("Mês com mais gastos", month_label(year, month), None, f"{_money_text(spent)} em saídas"))
    else:
        rows.append(("Mês com mais gastos", "—", None, "Nenhum gasto no período"))

    uncategorized = [t for t in totals.counted if t.category is None]
    value = sum((abs(t.amount) for t in uncategorized), Decimal(0))
    rows.append(
        (
            "Sem categoria",
            len(uncategorized),
            COUNT_FORMAT,
            f"{_money_text(value)} em transações que nenhuma regra reconheceu"
            if uncategorized
            else "Todas as transações estão categorizadas",
        )
    )
    return rows


# ---------- Aba Por mês ----------


def _monthly_sheet(sheet: Worksheet, totals: Totals):
    sheet.sheet_properties.tabColor = CYAN
    headers = ["Mês", "Entradas", "Saídas", "Saldo", "Saldo acumulado", "Economia", "Transações"]
    _set_widths(sheet, [14, 16, 16, 16, 18, 12, 13])
    _title(
        sheet,
        "Por mês",
        "Economia = quanto das entradas sobrou no mês (saldo ÷ entradas). "
        "Saldo acumulado = soma dos saldos até aquele mês.",
        columns=len(headers),
    )

    header_row = TABLE_START
    _header_row(sheet, header_row, headers)
    row = header_row + 1
    running = Decimal(0)
    for index, ((year, month), (income, spent, count)) in enumerate(totals.months.items()):
        balance = income - spent
        running += balance
        _write_row(
            sheet,
            row,
            [month_label(year, month), income, spent, balance, running, float(balance / income) if income else None, count],
            formats={2: MONEY_FORMAT, 3: MONEY_FORMAT, 4: MONEY_FORMAT, 5: MONEY_FORMAT, 6: PERCENT_FORMAT, 7: COUNT_FORMAT},
            zebra=index % 2 == 1,
        )
        sheet.cell(row=row, column=2).font = _font(color=GREEN)
        sheet.cell(row=row, column=3).font = _font(color=RED)
        row += 1

    last = row - 1
    _total_row(sheet, row, header_row, sum_columns=["B", "C", "D", "G"], formats={2: MONEY_FORMAT, 3: MONEY_FORMAT, 4: MONEY_FORMAT, 7: COUNT_FORMAT})
    # Economia do período todo: saldo total ÷ entradas totais
    economy = sheet.cell(row=row, column=6, value=f'=IF(B{row}>0,D{row}/B{row},"")')
    economy.number_format = PERCENT_FORMAT
    economy.font = _font(bold=True)
    economy.border = TOTAL_LINE
    sheet.cell(row=row, column=5).border = TOTAL_LINE

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if last > header_row:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Entradas x saídas por mês"
        chart.y_axis.title = "R$"
        chart.y_axis.numFmt = '"R$" #,##0'
        chart.y_axis.majorGridlines = None
        chart.height, chart.width = 8, 18
        data = Reference(sheet, min_col=2, max_col=3, min_row=header_row, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=header_row + 1, max_row=last))
        chart.series[0].graphicalProperties.solidFill = GREEN
        chart.series[1].graphicalProperties.solidFill = RED
        sheet.add_chart(chart, f"A{row + 3}")


# ---------- Aba Por categoria ----------


def _category_sheet(sheet: Worksheet, totals: Totals):
    sheet.sheet_properties.tabColor = GREEN
    _set_widths(sheet, [30, 16, 14, 13, 20])
    _title(
        sheet,
        "Por categoria",
        "Para onde foi o dinheiro (saídas) e de onde veio (entradas). "
        "% = fatia da categoria no total da tabela.",
        columns=5,
    )

    # Saídas
    header_row = TABLE_START
    _header_row(sheet, header_row, ["Categoria", "Saídas", "% das saídas", "Transações", "Média por transação"])
    row = _category_rows(sheet, header_row + 1, totals.spending, totals.expenses, value_color=RED, with_average=True)
    spending_last = row - 1
    _total_row(sheet, row, header_row, sum_columns=["B", "D"], formats={2: MONEY_FORMAT, 4: COUNT_FORMAT})
    share = sheet.cell(row=row, column=3, value=1 if totals.expenses else None)
    share.number_format, share.font, share.border = PERCENT_FORMAT, _font(bold=True), TOTAL_LINE
    sheet.cell(row=row, column=5).border = TOTAL_LINE
    if spending_last > header_row:
        # Barrinha dentro da célula: dá para comparar as fatias de olho
        sheet.conditional_formatting.add(
            f"C{header_row + 1}:C{spending_last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="F1948A"),
        )

    # Entradas
    income_header = row + 3
    sheet.cell(row=income_header - 1, column=1, value="Entradas por categoria").font = _font(bold=True, size=12, color=NAVY)
    _header_row(sheet, income_header, ["Categoria", "Entradas", "% das entradas", "Transações"])
    row = _category_rows(sheet, income_header + 1, totals.earning, totals.income, value_color=GREEN, with_average=False)
    _total_row(sheet, row, income_header, sum_columns=["B", "D"], formats={2: MONEY_FORMAT, 4: COUNT_FORMAT})
    share = sheet.cell(row=row, column=3, value=1 if totals.income else None)
    share.number_format, share.font, share.border = PERCENT_FORMAT, _font(bold=True), TOTAL_LINE

    if spending_last > header_row:
        chart = BarChart()
        chart.type = "bar"  # barras deitadas: nomes de categoria longos cabem
        chart.title = "Saídas por categoria"
        chart.legend = None
        chart.x_axis.scaling.orientation = "maxMin"  # maior gasto em cima, como na tabela
        chart.y_axis.numFmt = '"R$" #,##0'
        chart.y_axis.majorGridlines = None
        chart.height = max(7, 1 + 0.7 * (spending_last - header_row))
        chart.width = 16
        chart.add_data(Reference(sheet, min_col=2, min_row=header_row, max_row=spending_last), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=header_row + 1, max_row=spending_last))
        chart.series[0].graphicalProperties.solidFill = CYAN
        sheet.add_chart(chart, f"G{header_row}")

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.sheet_view.showGridLines = False


def _category_rows(sheet, row, values: dict, total: Decimal, value_color: str, with_average: bool) -> int:
    for index, (name, (value, count)) in enumerate(values.items()):
        cells = [name, value, float(value / total) if total else 0, count]
        formats = {2: MONEY_FORMAT, 3: PERCENT_FORMAT, 4: COUNT_FORMAT}
        if with_average:
            cells.append(value / count)
            formats[5] = MONEY_FORMAT
        _write_row(sheet, row, cells, formats=formats, zebra=index % 2 == 1)
        sheet.cell(row=row, column=2).font = _font(color=value_color)
        row += 1
    return row


# ---------- Aba Transações ----------


def _transactions_sheet(sheet: Worksheet, transactions):
    sheet.sheet_properties.tabColor = MUTED
    headers = ["Data", "Descrição", "Categoria", "Tipo", "Valor", "Extrato", "Nos totais"]
    _set_widths(sheet, [12, 46, 24, 10, 16, 28, 16])
    _title(
        sheet,
        "Transações",
        "Todos os lançamentos do período. Use as setas do cabeçalho para filtrar ou ordenar. "
        "Linhas em cinza são de categorias ignoradas nos totais.",
        columns=len(headers),
    )
    header_row = TABLE_START
    _header_row(sheet, header_row, headers)

    row = header_row + 1
    for index, t in enumerate(transactions):
        counted = counts_in_reports(t)
        _write_row(
            sheet,
            row,
            [
                t.date,
                t.description,
                _category_name(t),
                "Entrada" if t.amount > 0 else "Saída",
                t.amount,
                t.statement.filename,
                "Sim" if counted else "Não (ignorada)",
            ],
            formats={1: DATE_FORMAT, 5: SIGNED_MONEY_FORMAT},
            zebra=index % 2 == 1,
        )
        if not counted:
            for column in range(1, len(headers) + 1):
                sheet.cell(row=row, column=column).font = _font(italic=True, color="8A96A3")
        row += 1

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)  # cabeçalho fixo ao rolar
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(row - 1, header_row)}"
    # Impressão: repete o cabeçalho em cada página
    sheet.print_title_rows = f"{header_row}:{header_row}"


# ---------- Helpers ----------


def _period_label(filters: ReportFilters) -> str:
    start = filters.start_date.strftime("%d/%m/%Y") if filters.start_date else None
    end = filters.end_date.strftime("%d/%m/%Y") if filters.end_date else None
    if start and end:
        return f"{start} a {end}"
    if start:
        return f"A partir de {start}"
    if end:
        return f"Até {end}"
    return "Todo o período"


def _money_text(value: Decimal) -> str:
    """Decimal("1234.5") -> "R$ 1.234,50" (para textos explicativos)"""
    text = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {text}"


def _plural(count: int, singular: str, plural: str) -> str:
    return f"{count} {singular if count == 1 else plural}"


def _percent_text(value: Decimal) -> str:
    return f"{value * 100:.1f}%".replace(".", ",")


def _title(sheet: Worksheet, title: str, subtitle: str, columns: int):
    """Faixa escura com o título (linha 1) e a explicação da aba (linha 2)."""
    last = get_column_letter(columns)
    sheet.merge_cells(f"A1:{last}1")
    sheet.merge_cells(f"A2:{last}2")
    cell = sheet["A1"]
    cell.value = title
    cell.font = _font(bold=True, size=16, color="FFFFFF")
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 34
    note = sheet["A2"]
    note.value = subtitle
    note.font = _font(italic=True, color=MUTED)
    note.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    sheet.row_dimensions[2].height = 32


def _section(sheet: Worksheet, row: int, title: str) -> int:
    cell = sheet.cell(row=row, column=1, value=title)
    cell.font = _font(bold=True, size=12, color=NAVY)
    for column in range(1, 4):
        sheet.cell(row=row, column=column).border = Border(bottom=Side(style="medium", color=CYAN))
    sheet.row_dimensions[row].height = 22
    return row + 1


def _label_rows(sheet: Worksheet, row: int, rows: list[tuple], colors: Optional[dict] = None) -> int:
    """Linhas "rótulo | valor | o que significa" do Resumo."""
    for label, value, number_format, explanation in rows:
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = _font(bold=True)
        value_cell = sheet.cell(row=row, column=2, value=value)
        value_cell.font = _font(bold=True, color=(colors or {}).get(label, NAVY))
        value_cell.alignment = Alignment(horizontal="right", vertical="top")
        if number_format and not isinstance(value, str):
            value_cell.number_format = number_format
        note = sheet.cell(row=row, column=3, value=explanation)
        note.font = _font(color=MUTED)
        note.alignment = WRAP
        for column in range(1, 4):
            sheet.cell(row=row, column=column).border = BOTTOM_LINE
        label_cell.alignment = Alignment(vertical="top")
        row += 1
    return row


def _note(sheet: Worksheet, row: int, text: str, columns: int):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = _font(color=MUTED)
    cell.alignment = WRAP
    sheet.row_dimensions[row].height = 28


def _header_row(sheet: Worksheet, row: int, titles: list[str]) -> int:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left" if column == 1 else "center", vertical="center")
    sheet.row_dimensions[row].height = 22
    return row + 1


def _write_row(sheet: Worksheet, row: int, values: list, formats: dict[int, str], zebra: bool):
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = _font()
        cell.border = BOTTOM_LINE
        if zebra:
            cell.fill = ZEBRA_FILL
        if column in formats:
            cell.number_format = formats[column]


def _total_row(sheet: Worksheet, row: int, header_row: int, sum_columns: list[str], formats: dict[int, str]):
    """Linha de total com fórmula SUM, para continuar certa se editarem a planilha."""
    label = sheet.cell(row=row, column=1, value="Total")
    label.font = _font(bold=True)
    label.border = TOTAL_LINE
    first, last = header_row + 1, row - 1
    for column in sum_columns:
        cell = sheet[f"{column}{row}"]
        cell.value = f"=SUM({column}{first}:{column}{last})" if last >= first else 0
        cell.font = _font(bold=True)
        cell.border = TOTAL_LINE
        index = cell.column
        if index in formats:
            cell.number_format = formats[index]


def _set_widths(sheet: Worksheet, widths: list[int]):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _finish_sheet(sheet: Worksheet):
    """Impressão em paisagem, cabendo na largura da página."""
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
