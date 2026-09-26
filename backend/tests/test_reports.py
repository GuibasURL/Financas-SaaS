import io
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from tests.conftest import CSV_FEVEREIRO, CSV_JANEIRO, upload_csv

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _export(client, **params):
    return client.get("/reports/export", params=params)


def _workbook(response):
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def _table(sheet, header_row=4):
    """Linhas de uma tabela (sem o cabeçalho), até a primeira linha vazia."""
    rows = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0] is None:
            break
        rows.append(list(row))
    return rows


def _header(sheet, row=4):
    return [c.value for c in sheet[row] if c.value is not None]


def _summary(workbook) -> dict:
    """Resumo como {rótulo: valor} (coluna A -> coluna B)."""
    return {
        row[0]: row[1]
        for row in workbook["Resumo"].iter_rows(max_col=2, values_only=True)
        if row[0] and row[1] is not None
    }


def _notes(workbook) -> dict:
    """Resumo como {rótulo: explicação} (coluna A -> coluna C)."""
    return {
        row[0]: row[2]
        for row in workbook["Resumo"].iter_rows(max_col=3, values_only=True)
        if row[0] and row[2]
    }


def _labels(sheet) -> list:
    return [row[0] for row in sheet.iter_rows(max_col=1, values_only=True)]


@pytest.fixture
def dados(upload, categories):
    """
    Janeiro: IFOOD -45,90 (Alimentação), SALARIO +5000, UBER -20 (Transporte), FARMACIA -30,10
    Fevereiro: IFOOD -60 (Alimentação), UBER -15,50 (Transporte)
    """
    janeiro = upload(CSV_JANEIRO, filename="janeiro.csv").json()[0]["statement_id"]
    fevereiro = upload(CSV_FEVEREIRO, filename="fevereiro.csv").json()[0]["statement_id"]
    return {"janeiro": janeiro, "fevereiro": fevereiro}


# ---------- Arquivo ----------


def test_devolve_xlsx_com_as_quatro_abas(client, dados):
    response = _export(client)

    assert response.headers["content-type"] == XLSX
    assert _workbook(response).sheetnames == ["Resumo", "Por mês", "Por categoria", "Transações"]


@pytest.mark.parametrize(
    "params, filename",
    [
        ({}, "vexira-relatorio_completo.xlsx"),
        ({"start_date": "2025-01-01", "end_date": "2025-01-31"}, "vexira-relatorio_2025-01-01_a_2025-01-31.xlsx"),
        ({"start_date": "2025-02-01"}, "vexira-relatorio_2025-02-01_a_hoje.xlsx"),
        ({"end_date": "2025-01-31"}, "vexira-relatorio_inicio_a_2025-01-31.xlsx"),
    ],
)
def test_nome_do_arquivo_reflete_o_periodo(client, dados, params, filename):
    response = _export(client, **params)

    assert response.headers["content-disposition"] == f'attachment; filename="{filename}"'


def test_nome_do_arquivo_fica_visivel_para_o_frontend(client):
    response = client.get("/", headers={"Origin": "http://localhost:5173"})

    assert "content-disposition" in response.headers["access-control-expose-headers"].lower()


# ---------- Conteúdo ----------


def test_resumo_com_totais_e_filtros(client, dados):
    workbook = _workbook(_export(client))
    summary = _summary(workbook)

    assert workbook["Resumo"]["A1"].value == "Vexira — Relatório financeiro"
    assert summary["Período"] == "Todo o período"
    assert summary["Extrato"] == "Todos"
    assert summary["Entradas"] == pytest.approx(5000.00)
    # Saídas em valor positivo: "quanto saiu"
    assert summary["Saídas"] == pytest.approx(171.50)
    assert summary["Saldo"] == pytest.approx(4828.50)
    assert summary["Taxa de economia"] == pytest.approx(4828.50 / 5000)
    assert summary["Média de saídas por mês"] == pytest.approx(171.50 / 2)
    assert summary["Transações"] == 6
    datetime.strptime(summary["Gerado em"], "%d/%m/%Y %H:%M")


def test_resumo_explica_cada_numero(client, dados):
    workbook = _workbook(_export(client))
    notes = _notes(workbook)

    assert "Entradas menos saídas" in notes["Saldo"]
    assert "saldo ÷ entradas" in notes["Taxa de economia"]
    assert "2 meses" in notes["Média de saídas por mês"]
    assert "Como ler este relatório" in _labels(workbook["Resumo"])


def test_destaques(client, dados):
    workbook = _workbook(_export(client))
    summary, notes = _summary(workbook), _notes(workbook)

    assert summary["Maior gasto"] == pytest.approx(60.00)
    assert notes["Maior gasto"] == "IFOOD *PIZZARIA em 03/02/2025"
    assert summary["Maior entrada"] == pytest.approx(5000.00)
    assert summary["Categoria que mais pesou"] == "Alimentação"
    assert notes["Categoria que mais pesou"] == "R$ 105,90 em 2 transações: 61,7% das saídas"
    assert summary["Mês com mais gastos"] == "jan/2025"
    assert notes["Mês com mais gastos"] == "R$ 96,00 em saídas"
    # FARMACIA (saída) e SALARIO (entrada): nenhuma categoria do teste reconhece
    assert summary["Sem categoria"] == 2
    assert notes["Sem categoria"] == "R$ 5.030,10 em transações que nenhuma regra reconheceu"


def test_cinco_maiores_gastos(client, dados):
    sheet = _workbook(_export(client))["Resumo"]
    header = _labels(sheet).index("Descrição") + 1  # linha do cabeçalho (1-based)

    rows = _table(sheet, header_row=header)

    assert [(r[0], r[1]) for r in rows] == [
        ("IFOOD *PIZZARIA", pytest.approx(60.00)),
        ("IFOOD *RESTAURANTE XYZ", pytest.approx(45.90)),
        ("FARMACIA", pytest.approx(30.10)),
        ("UBER TRIP", pytest.approx(20.00)),
        ("UBER TRIP", pytest.approx(15.50)),
    ]
    assert rows[0][2] == "Alimentação · 03/02/2025"


def test_gerado_em_usa_o_fuso_do_app_e_nao_o_do_servidor(client, dados, monkeypatch):
    # Um fuso bem longe do servidor e de Brasília: se a hora vier certa, veio do APP_TIMEZONE
    tokyo = ZoneInfo("Asia/Tokyo")
    monkeypatch.setattr("app.routers.reports.APP_TIMEZONE", tokyo)

    summary = _summary(_workbook(_export(client)))

    generated = datetime.strptime(summary["Gerado em"], "%d/%m/%Y %H:%M")
    now = datetime.now(tokyo).replace(tzinfo=None)
    assert now - timedelta(minutes=2) <= generated <= now


def test_por_mes(client, dados):
    sheet = _workbook(_export(client))["Por mês"]
    rows = _table(sheet)

    assert _header(sheet) == ["Mês", "Entradas", "Saídas", "Saldo", "Saldo acumulado", "Economia", "Transações"]
    assert rows[0] == [
        "jan/2025", 5000, pytest.approx(96.00), pytest.approx(4904.00), pytest.approx(4904.00),
        pytest.approx(4904 / 5000), 4,
    ]
    # Sem entradas no mês: economia fica em branco (não divide por zero)
    assert rows[1] == ["fev/2025", 0, pytest.approx(75.50), pytest.approx(-75.50), pytest.approx(4828.50), None, 2]
    # Linha de total com fórmula, para continuar certa se editarem a planilha
    assert rows[2] == ["Total", "=SUM(B5:B6)", "=SUM(C5:C6)", "=SUM(D5:D6)", None, '=IF(B7>0,D7/B7,"")', "=SUM(G5:G6)"]
    assert sheet.freeze_panes == "A5"
    assert len(sheet._charts) == 1


def test_por_categoria_inclui_sem_categoria_e_ordena_pelo_maior_gasto(client, dados):
    sheet = _workbook(_export(client))["Por categoria"]
    rows = _table(sheet)

    assert _header(sheet) == ["Categoria", "Saídas", "% das saídas", "Transações", "Média por transação"]
    assert [r[0] for r in rows] == ["Alimentação", "Transporte", "Sem categoria", "Total"]
    assert rows[0][1:] == [pytest.approx(105.90), pytest.approx(105.90 / 171.50), 2, pytest.approx(52.95)]
    assert rows[2][1] == pytest.approx(30.10)  # FARMACIA
    assert rows[3][1:4] == ["=SUM(B5:B7)", 1, "=SUM(D5:D7)"]
    assert len(sheet._charts) == 1


def test_por_categoria_tambem_mostra_as_entradas(client, dados):
    sheet = _workbook(_export(client))["Por categoria"]
    header = _labels(sheet).index("Entradas por categoria") + 2

    assert _header(sheet, header) == ["Categoria", "Entradas", "% das entradas", "Transações"]
    first = header + 1
    assert _table(sheet, header) == [
        ["Sem categoria", 5000, 1, 1, None],  # SALARIO: nenhuma categoria do teste reconhece
        ["Total", f"=SUM(B{first}:B{first})", 1, f"=SUM(D{first}:D{first})", None],
    ]


def test_transacoes_em_ordem_de_data_com_categoria_e_extrato(client, dados):
    sheet = _workbook(_export(client))["Transações"]
    rows = _table(sheet)

    assert _header(sheet) == ["Data", "Descrição", "Categoria", "Tipo", "Valor", "Extrato", "Nos totais"]
    assert len(rows) == 6
    assert [r[0].date().isoformat() for r in rows] == sorted(r[0].date().isoformat() for r in rows)
    first = rows[0]
    assert first[1:] == ["IFOOD *RESTAURANTE XYZ", "Alimentação", "Saída", pytest.approx(-45.90), "janeiro.csv", "Sim"]
    assert [r[3] for r in rows if r[1] == "SALARIO EMPRESA"] == ["Entrada"]
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref == "A4:G10"


def test_valores_e_datas_sao_numeros_formatados_nao_texto(client, dados):
    sheet = _workbook(_export(client))["Transações"]

    assert isinstance(sheet["A5"].value, datetime)
    assert sheet["A5"].number_format == "DD/MM/YYYY"
    assert isinstance(sheet["E5"].value, (int, float))
    assert "R$" in sheet["E5"].number_format


def test_visual_padronizado(client, dados):
    workbook = _workbook(_export(client))

    for sheet in workbook.worksheets:
        # Faixa de título escura e a mesma fonte em todas as abas
        assert sheet["A1"].fill.fgColor.rgb.endswith("0B1F33")
        assert sheet["A1"].font.name == "Arial"
        assert sheet.page_setup.orientation == "landscape"
    transactions = workbook["Transações"]
    # Linhas alternadas (zebra) na lista
    assert transactions["A5"].fill.fill_type is None
    assert transactions["A6"].fill.fgColor.rgb.endswith("F3F7FA")


def test_categoria_ignorada_fica_fora_dos_totais_mas_na_lista(client, dados, categories):
    client.patch(f"/categories/{categories['transporte'].id}", json={"ignore_in_reports": True})

    workbook = _workbook(_export(client))

    summary = _summary(workbook)
    assert summary["Saídas"] == pytest.approx(171.50 - 35.50)
    assert summary["Transações"] == 6
    assert summary["Fora dos totais"] == 2
    assert _notes(workbook)["Fora dos totais"].startswith("R$ 35,50 em categorias")
    assert "Transporte" not in [r[0] for r in _table(workbook["Por categoria"])]
    sheet = workbook["Transações"]
    uber = [(i, r) for i, r in enumerate(_table(sheet), start=5) if r[1] == "UBER TRIP"]
    assert [r[6] for _, r in uber] == ["Não (fora dos totais)", "Não (fora dos totais)"]
    # Em cinza e itálico, para se destacar das que contam
    assert all(sheet.cell(row=i, column=2).font.i for i, _ in uber)


def test_so_sem_categoria_nao_vira_a_categoria_que_mais_pesou(client):
    upload_csv(client, "data,descricao,valor\n2025-01-05,LOJA X,-10.00\n")

    summary = _summary(_workbook(_export(client)))

    assert summary["Categoria que mais pesou"] == "—"
    assert summary["Sem categoria"] == 1


# ---------- Filtros ----------


def test_filtro_de_periodo_inclui_as_datas_das_pontas(client, dados):
    workbook = _workbook(_export(client, start_date="2025-01-06", end_date="2025-02-03"))

    descriptions = [r[1] for r in _table(workbook["Transações"])]
    # 05/01 fica de fora; 06/01 e 03/02 entram (pontas inclusivas); 20/02 fica de fora
    assert descriptions == ["SALARIO EMPRESA", "UBER TRIP", "FARMACIA", "IFOOD *PIZZARIA"]
    assert _summary(workbook)["Período"] == "06/01/2025 a 03/02/2025"


def test_filtro_de_extrato(client, dados):
    workbook = _workbook(_export(client, statement_id=dados["fevereiro"]))

    assert len(_table(workbook["Transações"])) == 2
    assert _summary(workbook)["Extrato"] == "fevereiro.csv"


def test_periodo_sem_transacoes_gera_planilha_vazia(client, dados):
    workbook = _workbook(_export(client, start_date="2030-01-01"))

    summary = _summary(workbook)
    assert summary["Transações"] == 0
    assert summary["Taxa de economia"] == "—"
    assert summary["Maior gasto"] == "—"
    assert summary["Maior entrada"] == "—"
    assert summary["Mês com mais gastos"] == "—"
    assert _notes(workbook)["Sem categoria"] == "Todas as transações estão categorizadas"
    assert _table(workbook["Transações"]) == []
    assert _table(workbook["Por mês"]) == [["Total", 0, 0, 0, None, '=IF(B5>0,D5/B5,"")', 0]]
    assert _table(workbook["Por categoria"]) == [["Total", 0, None, 0, None]]
    # Sem dados, sem gráfico vazio
    assert workbook["Por mês"]._charts == [] and workbook["Por categoria"]._charts == []


def test_data_inicial_depois_da_final_da_400(client):
    response = _export(client, start_date="2025-02-01", end_date="2025-01-01")

    assert response.status_code == 400


def test_data_invalida_da_422(client):
    assert _export(client, start_date="01/02/2025").status_code == 422


def test_extrato_inexistente_da_404(client):
    assert _export(client, statement_id=999).status_code == 404


# ---------- Autenticação e isolamento ----------


def test_exige_token(anon_client):
    assert anon_client.get("/reports/export").status_code == 401


def test_relatorio_so_tem_dados_do_proprio_usuario(client, other_client, dados):
    upload_csv(other_client, CSV_FEVEREIRO, filename="do-bruno.csv")

    workbook = _workbook(_export(other_client))

    rows = _table(workbook["Transações"])
    assert len(rows) == 2
    assert {r[5] for r in rows} == {"do-bruno.csv"}


def test_extrato_de_outro_usuario_da_404(other_client, dados):
    assert _export(other_client, statement_id=dados["janeiro"]).status_code == 404
