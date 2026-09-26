"""
Segurança, parte 2: ataques que não são "B mexendo nos dados de A".

- Senha gigante (acima dos 72 bytes do bcrypt) não derruba a API
- Token sem validade, com outro algoritmo, com dono inválido ou de uma
  conta já excluída (cujo id foi reaproveitado) não entra
- Trocar o e-mail cancela o link de "Esqueceu a senha?" do e-mail antigo
- Campos extras no corpo (user_id, id, hashed_password) são ignorados
- CORS, SQL injection, fórmula no nome da categoria e do arquivo
- Limites de tamanho do que fica guardado no banco
- Limite de cadastros por IP (contas em massa e teste de quais e-mails existem)
"""
import io
from datetime import datetime, timedelta, timezone

import jwt
import pytest
from openpyxl import load_workbook

from app import config
from app.config import CORS_ORIGINS, SECRET_KEY
from app.models.category import Category
from app.models.user import User
from app.routers.auth import register_limiter
from app.services.password_reset import create_reset_token
from app.services.security import create_access_token
from tests.conftest import CSV_JANEIRO, PASSWORD, TEST_DATABASE_URL, _create_user, upload_csv

NEW_PASSWORD = "Girassol#Azul91"


def _me(anon_client, token: str):
    return anon_client.get("/auth/me", headers={"Authorization": f"Bearer {token}"})


def _login(anon_client, email: str, password: str):
    return anon_client.post("/auth/login", data={"username": email, "password": password})


# ---------- Senha gigante (mais de 72 bytes, o limite do bcrypt) ----------

HUGE_PASSWORD = "b" * 1000


def test_senha_gigante_no_login_e_senha_errada_e_nao_erro(anon_client, user):
    response = _login(anon_client, user.email, HUGE_PASSWORD)

    assert (response.status_code, response.json()["detail"]) == (401, "E-mail ou senha incorretos")


def test_senha_gigante_conta_no_limite_de_tentativas(anon_client, user):
    for _ in range(5):
        _login(anon_client, user.email, HUGE_PASSWORD)

    assert _login(anon_client, user.email, PASSWORD).status_code == 429


def test_senha_certa_mais_lixo_depois_dos_72_bytes_nao_entra(anon_client, user):
    # O bcrypt antigo cortava a senha em 72 bytes: "senha certa + lixo" entraria
    assert _login(anon_client, user.email, PASSWORD + "x" * 100).status_code == 401


@pytest.mark.parametrize(
    "method, path, body",
    [
        ("POST", "/auth/me/password", {"current_password": HUGE_PASSWORD, "new_password": NEW_PASSWORD}),
        ("PATCH", "/auth/me", {"name": "Ana", "email": "outro@teste.com", "current_password": HUGE_PASSWORD}),
        ("DELETE", "/auth/me", {"password": HUGE_PASSWORD}),
    ],
)
def test_senha_atual_gigante_e_senha_errada_e_nao_erro(client, method, path, body):
    response = client.request(method, path, json=body)

    assert (response.status_code, response.json()["detail"]) == (400, "Senha atual incorreta")
    assert client.get("/auth/me").json()["email"] == "ana@teste.com"


# ---------- Tokens: validade, algoritmo e conta dona ----------


def _signed(payload: dict, algorithm: str = "HS256") -> str:
    return jwt.encode(payload, SECRET_KEY, algorithm=algorithm)


def _in_one_hour() -> datetime:
    return datetime.now(timezone.utc) + timedelta(hours=1)


def test_token_sem_validade_nao_entra(anon_client, user):
    # Mesmo assinado com a chave certa: sem "exp", valeria para sempre
    assert _me(anon_client, _signed({"sub": str(user.id)})).status_code == 401


@pytest.mark.filterwarnings("ignore::jwt.warnings.InsecureKeyLengthWarning")
def test_token_com_outro_algoritmo_nao_entra(anon_client, user):
    token = _signed({"sub": str(user.id), "exp": _in_one_hour()}, algorithm="HS512")

    assert _me(anon_client, token).status_code == 401


@pytest.mark.parametrize("sub", ["abc", "1 OR 1=1", "", None])
def test_token_com_dono_invalido_nao_entra(anon_client, user, sub):
    payload = {"exp": _in_one_hour()} if sub is None else {"sub": sub, "exp": _in_one_hour()}

    assert _me(anon_client, _signed(payload)).status_code == 401


def test_token_de_conta_excluida_nao_abre_conta_nova_com_o_mesmo_id(anon_client, db_session, user):
    # Conta da Ana criada ontem, token emitido há uma hora (ainda válido)
    user.created_at = datetime.now(timezone.utc) - timedelta(days=1)
    db_session.commit()
    old_token = create_access_token(user.id, issued_at=datetime.now(timezone.utc) - timedelta(hours=1))
    old_id = user.id
    headers = {"Authorization": f"Bearer {old_token}"}
    assert anon_client.request("DELETE", "/auth/me", json={"password": PASSWORD}, headers=headers).status_code == 204

    newcomer = _create_user(db_session, "carla@teste.com")
    if TEST_DATABASE_URL is None:
        # O SQLite reaproveita o id da conta excluída (o Postgres, não)
        assert newcomer.id == old_id

    assert _me(anon_client, old_token).status_code == 401
    # A dona da conta nova entra normalmente
    assert _me(anon_client, create_access_token(newcomer.id)).json()["email"] == "carla@teste.com"


# ---------- Link de "Esqueceu a senha?" x troca de e-mail ----------


def _reset(anon_client, raw_token: str):
    return anon_client.post("/auth/reset-password", json={"token": raw_token, "new_password": NEW_PASSWORD})


def test_trocar_o_email_cancela_o_link_mandado_para_o_email_antigo(client, anon_client, db_session, user):
    # Quem invadiu o e-mail antigo não pode usar um link que chegou lá
    raw_token = create_reset_token(db_session, user)
    db_session.commit()
    client.patch("/auth/me", json={"name": "Ana", "email": "ana.nova@teste.com", "current_password": PASSWORD})

    assert _reset(anon_client, raw_token).status_code == 400
    assert _login(anon_client, "ana.nova@teste.com", PASSWORD).status_code == 200


def test_salvar_o_perfil_sem_trocar_o_email_nao_cancela_o_link(client, anon_client, db_session, user):
    raw_token = create_reset_token(db_session, user)
    db_session.commit()
    client.patch("/auth/me", json={"name": "Ana Maria", "email": "ana@teste.com"})

    assert _reset(anon_client, raw_token).status_code == 200


# ---------- Campos extras: ninguém cria dados em nome de outro ----------


def test_categoria_com_user_id_de_outro_fica_com_quem_criou(client, other_client, user, db_session):
    other_id = other_client.get("/auth/me").json()["id"]

    response = client.post("/categories", json={"name": "Invasora", "user_id": other_id, "id": 424242})

    assert response.json()["id"] != 424242
    assert db_session.get(Category, response.json()["id"]).user_id == user.id
    assert other_client.get("/categories").json() == []


def test_cadastro_ignora_campos_extras(anon_client, db_session):
    response = anon_client.post(
        "/auth/register",
        json={"email": "nova@teste.com", "password": NEW_PASSWORD, "id": 999, "hashed_password": "x"},
    )

    user = db_session.query(User).filter(User.email == "nova@teste.com").one()
    assert response.json()["id"] == user.id != 999
    assert user.hashed_password.startswith("$2b$")
    assert _login(anon_client, "nova@teste.com", NEW_PASSWORD).status_code == 200


def test_foto_vale_pelo_conteudo_e_nao_pelo_tipo_que_o_cliente_diz(client):
    # Diz ser SVG (que pode ter script), mas o conteúdo é PNG: vira PNG
    png = b"\x89PNG\r\n\x1a\n<svg onload=alert(1)>"
    response = client.put("/auth/me/avatar", files={"file": ("a.svg", png, "image/svg+xml")})
    assert response.json()["avatar_url"].startswith("data:image/png;base64,")

    # E um SVG de verdade, dizendo ser PNG, é recusado
    svg = client.put("/auth/me/avatar", files={"file": ("a.png", b"<svg onload=alert(1)>", "image/png")})
    assert svg.status_code == 400


# ---------- CORS: só o frontend configurado chama a API ----------

EVIL_ORIGIN = "https://site-malicioso.exemplo"


def test_cors_nao_libera_outro_site(anon_client):
    preflight = anon_client.options(
        "/auth/login", headers={"Origin": EVIL_ORIGIN, "Access-Control-Request-Method": "POST"}
    )
    response = anon_client.get("/", headers={"Origin": EVIL_ORIGIN})

    assert "access-control-allow-origin" not in preflight.headers
    assert "access-control-allow-origin" not in response.headers


def test_cors_libera_o_frontend(anon_client):
    response = anon_client.get("/", headers={"Origin": CORS_ORIGINS[0]})

    assert response.headers["access-control-allow-origin"] == CORS_ORIGINS[0]


# ---------- SQL injection ----------


@pytest.mark.parametrize("params", [{"month": "1 OR 1=1"}, {"statement_id": "1; DROP TABLE users"}])
def test_filtro_com_sql_e_recusado(client, params):
    assert client.get("/transactions", params=params).status_code == 422


def test_categoria_com_sql_e_curinga_e_so_texto(client, other_client):
    upload_csv(other_client, CSV_JANEIRO)  # dados do Bruno, que um "%" não pode alcançar
    upload_csv(client, CSV_JANEIRO)

    response = client.post("/categories", json={"name": "'; DROP TABLE users;--", "keywords": "%,_,' or '1'='1"})

    assert response.status_code == 200
    assert client.post("/categories/apply-rules").json() == {"categorized": 0}
    assert client.get("/auth/me").status_code == 200
    assert all(t["category_id"] is None for t in other_client.get("/transactions").json())


# ---------- Planilha: nome de categoria e de arquivo com fórmula ----------


def test_nome_de_categoria_e_de_arquivo_com_formula_viram_texto(client):
    client.post("/categories", json={"name": '=HYPERLINK("http://atacante.exemplo","clique")', "keywords": "ifood"})
    filename = "=cmd|' /C calc'!A0.csv"
    statement_id = upload_csv(client, CSV_JANEIRO, filename=filename).json()[0]["statement_id"]

    report = client.get("/reports/export", params={"statement_id": statement_id})
    workbook = load_workbook(io.BytesIO(report.content))

    cells = [
        cell
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and ("HYPERLINK" in cell.value or "cmd|" in cell.value)
    ]
    # A categoria e o nome do extrato aparecem na planilha...
    assert any("HYPERLINK" in cell.value for cell in cells)
    assert any("cmd|" in cell.value for cell in cells)
    # ...sempre como texto, nunca como fórmula
    assert all(cell.data_type == "s" for cell in cells)


# ---------- Limites de tamanho (contra encher o banco) ----------


def test_nome_e_palavras_chave_da_categoria_tem_limite(client):
    assert client.post("/categories", json={"name": "N" * 101}).status_code == 422
    assert client.post("/categories", json={"name": "Ok", "keywords": "k" * 2001}).status_code == 422
    assert client.post("/categories", json={"name": "N" * 100, "keywords": "k" * 2000}).status_code == 200

    category_id = client.get("/categories").json()[0]["id"]
    assert client.patch(f"/categories/{category_id}", json={"name": "N" * 101}).status_code == 422
    assert client.patch(f"/categories/{category_id}", json={"keywords": "k" * 2001}).status_code == 422
    category = client.get("/categories").json()[0]
    assert (category["name"], category["keywords"]) == ("N" * 100, "k" * 2000)


def test_nome_do_arquivo_tem_limite(client):
    response = upload_csv(client, CSV_JANEIRO, filename="a" * 252 + ".csv")

    assert (response.status_code, response.json()["detail"]) == (
        400,
        "O nome do arquivo pode ter no máximo 255 caracteres",
    )
    assert client.get("/statements").json() == []
    assert upload_csv(client, CSV_JANEIRO, filename="a" * 251 + ".csv").status_code == 200


# ---------- Limite de cadastros por IP ----------


@pytest.fixture
def register_limit(monkeypatch):
    """Limite baixo (3) para o teste não precisar de muitos cadastros."""
    monkeypatch.setattr(register_limiter, "max_per_ip", 3)
    monkeypatch.setattr(register_limiter, "max_per_account", 3)


def _register(anon_client, email: str, **kwargs):
    return anon_client.post("/auth/register", json={"email": email, "password": NEW_PASSWORD}, **kwargs)


def test_limite_de_cadastros_por_ip(anon_client, db_session, register_limit):
    assert [_register(anon_client, f"conta{i}@teste.com").status_code for i in range(3)] == [201, 201, 201]

    response = _register(anon_client, "conta3@teste.com")

    assert response.status_code == 429
    assert response.json()["detail"] == "Muitos cadastros a partir desta conexão. Tente de novo em 60 minutos."
    assert int(response.headers["retry-after"]) > 59 * 60
    assert db_session.query(User).filter(User.email == "conta3@teste.com").first() is None


def test_testar_emails_no_cadastro_tambem_conta(anon_client, user, register_limit):
    # "E-mail já cadastrado" revela quem tem conta: quem testa vários também para
    assert [_register(anon_client, user.email).status_code for _ in range(3)] == [400, 400, 400]

    assert _register(anon_client, "ninguem@teste.com").status_code == 429


def test_cadastro_invalido_nao_gasta_o_limite(anon_client, register_limit):
    # Senha fraca volta 422 antes de chegar à rota: não revela nada nem cria conta
    for _ in range(5):
        anon_client.post("/auth/register", json={"email": "nova@teste.com", "password": "123"})

    assert _register(anon_client, "nova@teste.com").status_code == 201


def test_limite_de_cadastro_e_separado_do_login(anon_client, user, register_limit):
    for i in range(3):
        _register(anon_client, f"conta{i}@teste.com")

    assert _login(anon_client, user.email, PASSWORD).status_code == 200


def test_trocar_o_ip_no_header_nao_escapa_do_limite_de_cadastro(anon_client, register_limit, monkeypatch):
    monkeypatch.setattr(config, "TRUSTED_PROXY_HOPS", 1)

    def attempt(i: int, real_ip: str = "200.1.1.1"):
        # O visitante inventa um IP a cada tentativa; o proxy acrescenta o real
        return _register(anon_client, f"conta{i}@teste.com", headers={"X-Forwarded-For": f"6.6.6.{i}, {real_ip}"})

    assert [attempt(i).status_code for i in range(3)] == [201, 201, 201]
    assert attempt(3).status_code == 429
    # Outra pessoa, de outro IP de verdade, cadastra normalmente
    assert attempt(4, real_ip="177.2.2.2").status_code == 201
